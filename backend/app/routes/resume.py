from fastapi import FastAPI, APIRouter
from fastapi.responses import FileResponse
from pydantic import BaseModel
import docx
from docx.shared import Inches, Pt, RGBColor
import os
import tempfile

router = APIRouter()

class ResumeGenerationSchema(BaseModel):
    template_style: str
    file_format: str
    resume_data: dict

@router.post("/api/resume/generate")
def generate_resume_file(payload: ResumeGenerationSchema):
    data = payload.resume_data
    personal = data.get("personalInfo", {})
    style = payload.template_style
    
    doc = docx.Document()

    # Apply Style Themes
    primary_color = RGBColor(37, 99, 235) if style == "modern" else RGBColor(55, 65, 81) if style == "classic" else RGBColor(5, 150, 105)

    # 1. Header (Name & Contact)
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(personal.get("fullName", "Candidate Name"))
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = primary_color

    contact_p = doc.add_paragraph()
    contact_text = f"{personal.get('email', '')} | {personal.get('phone', '')} | {personal.get('location', '')}"
    if personal.get('linkedin'):
        contact_text += f" | LinkedIn: {personal.get('linkedin')}"
    contact_p.add_run(contact_text).font.size = Pt(10)

    # Summary
    if personal.get("summary"):
        doc.add_heading("Professional Summary", level=2)
        doc.add_paragraph(personal.get("summary"))

    # Tech Stack & Languages
    if data.get("techStack") or data.get("languages"):
        doc.add_heading("Skills & Languages", level=2)
        if data.get("techStack"):
            p = doc.add_paragraph()
            p.add_run("Tech Stack: ").bold = True
            p.add_run(data.get("techStack"))
        if data.get("languages"):
            p = doc.add_paragraph()
            p.add_run("Languages: ").bold = True
            p.add_run(data.get("languages"))

    # Education & Diplomas
    if data.get("education"):
        doc.add_heading("Education & Diplomas", level=2)
        for edu in data["education"]:
            p = doc.add_paragraph()
            p.add_run(f"{edu.get('degree')} - {edu.get('institution')} ").bold = True
            p.add_run(f"({edu.get('startYear')} - {edu.get('endYear')})\n")
            if edu.get('fieldOfStudy'):
                p.add_run(f"Stream: {edu.get('fieldOfStudy')} ")
            if edu.get('grade'):
                p.add_run(f"| Grade/Score: {edu.get('grade')}")

    # Experience
    if data.get("experience"):
        doc.add_heading("Experience", level=2)
        for exp in data["experience"]:
            p = doc.add_paragraph()
            p.add_run(f"{exp.get('role')} at {exp.get('company')} ").bold = True
            p.add_run(f"({exp.get('startDate')} - {exp.get('endDate')})\n")
            p.add_run(exp.get('description', ''))

    # Internships
    if data.get("internships"):
        doc.add_heading("Internships", level=2)
        for intern in data["internships"]:
            p = doc.add_paragraph()
            p.add_run(f"{intern.get('role')} - {intern.get('company')} ").bold = True
            p.add_run(f"({intern.get('duration')})\n")
            p.add_run(intern.get('description', ''))

    # Projects
    if data.get("projects"):
        doc.add_heading("Projects", level=2)
        for proj in data["projects"]:
            p = doc.add_paragraph()
            p.add_run(f"{proj.get('title')} ").bold = True
            if proj.get('techUsed'):
                p.add_run(f"[{proj.get('techUsed')}]\n")
            p.add_run(proj.get('description', ''))

    # Save to Temp File and Return as Download
    temp_dir = tempfile.gettempdir()
    file_path = os.path.join(temp_dir, f"resume_{payload.template_style}.docx")
    doc.save(file_path)

    return FileResponse(
        path=file_path,
        filename=f"{personal.get('fullName', 'Resume')}.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )